from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph


THESIS_PATH = Path(r"C:\Users\21080\Desktop\毕设\102201204_王碧强_基于LLM的简历智能生成与优化系统.docx")
BACKUP_PATH = THESIS_PATH.with_name(THESIS_PATH.stem + "_第5章第6章修改前备份.docx")


CH5_REWRITES = {
    "5.1 系统实现环境": [
        (
            "paragraph",
            "当前版本系统采用 Python 3.11 作为主要开发语言，前端与业务编排基于 Streamlit 完成，数据库层使用 SQLite，便于在单机环境下快速部署和调试。与早期仅强调“能够跑通”不同，现阶段实现更注重展示稳定性与结构清晰度：页面入口已收敛为薄入口 `web_app.py`，导航、会话状态和页面分发集中到 `wbq/ui/app.py`，用户中心、数据输入、模型配置、生成页、结果页分别拆分到 `wbq/ui/pages/`，页面内部的细粒度逻辑再下沉至 `wbq/ui/components/`。这种组织方式使界面代码、业务逻辑和工具函数的边界更加明确，后续维护成本也更低。",
        ),
        (
            "paragraph",
            "在运行依赖方面，系统集成了多类组件：大模型调用侧支持 Qwen、GPT、Gemini 与 Ollama；文档处理侧使用 PyPDF2、PyMuPDF 和 pdf2docx 完成文本提取与格式转换；渲染输出侧使用 Jinja2、LaTeX 以及 xelatex 生成简历与求职信 PDF；结果分析侧结合 jieba、scikit-learn、matplotlib 和 wordcloud 完成匹配度计算与可视化。考虑到毕业设计答辩通常在单机环境进行，系统对本地字体、LaTeX 编译路径、历史数据库字段兼容和演示数据载入都做了专门处理，因此在真实使用中既能覆盖完整流程，也能在部分依赖缺失时给出明确提示。",
        ),
        (
            "paragraph",
            "从部署方式看，当前工程仍保持轻量化特征：开发阶段通过虚拟环境启动 `streamlit run web_app.py` 即可运行；数据库默认落地为 `resume_system.db`；上传文件、模板资源和输出结果分别放在 `uploads/`、`wbq/templates/` 和 `output/` 目录中。对于答辩展示这一特定场景，这种“本地可独立运行”的实现方式更容易复现，也更方便在无外部服务依赖的情况下完成功能演示。",
        ),
    ],
    "5.2 用户管理与数据输入实现": [
        (
            "paragraph",
            "用户管理与数据输入是整套系统中最贴近实际操作的一层。当前版本在登录界面上采用了“用户名+密码”的前端展示形式，但为了保持轻量实现，后端仍以用户名作为会话建立的关键标识：页面读取用户名后调用数据库管理器创建或复用用户ID，再将 `user_id`、`username` 与导航状态写入 Streamlit 会话。这样做虽然不涉及真正的口令校验，但足以支撑单用户演示、多用户记录隔离以及历史数据回溯等需求。相比早期将所有逻辑堆在一个页面脚本中的做法，现在的用户中心还承担了工作台状态卡、历史概览和历史记录恢复三个角色，页面职责比以前清楚得多。",
        ),
        (
            "code_caption",
            "代码清单5-1 用户登录与会话恢复核心代码",
        ),
        (
            "code",
            "if st.button(\"进入系统 / 注册\", use_container_width=True, type=\"primary\"):\n"
            "    if not username_input:\n"
            "        st.warning(\"请输入用户名。\")\n"
            "    else:\n"
            "        _clear_user_work_state(include_identity=False)\n"
            "        st.session_state[\"user_id\"] = db.create_user(username_input)\n"
            "        st.session_state[\"username\"] = username_input\n"
            "        st.rerun()\n"
            "\n"
            "if load_col.button(\"加载此版本\", key=f\"load_hist_{item['id']}\"):\n"
            "    _load_history_record(item[\"id\"])",
        ),
        (
            "paragraph",
            "数据输入部分围绕“岗位描述输入”和“简历输入”两条主线展开。岗位描述支持文本粘贴与招聘链接抓取两种模式，并增加了职位标签、技术栈标签和语音输入辅助，方便用户在没有现成JD文本时快速组织内容。简历输入则同时支持 PDF/JSON 上传和手动表单填写：前者更贴近实际求职场景，后者则适合在答辩环境下演示字段级控制与结构化录入过程。值得注意的是，页面还提供了“一键载入 Demo 数据”能力，用于在网络、模型或现场时间受限时快速跑通一整套流程，这一点在系统展示中非常实用。",
        ),
        (
            "paragraph",
            "为了让后续模块尽可能少地关心“输入从哪里来”，系统会在这一层就把所有输入统一整理为标准化会话状态。例如，文本型 JD 最终会写入 `jd_text`，链接型 JD 会写入 `jd_url`；文件型简历会记录到 `resume_source`，手动表单则会被标准化为结构化 `form_data` 和 `manual_resume_form`。这种处理方式使生成页无需区分上游到底是 PDF、JSON 还是手填表单，只需要从会话中读取规范化数据即可。",
        ),
    ],
    "5.4 LaTeX模板渲染与PDF生成实现": [
        (
            "paragraph",
            "简历输出采用“结构化 JSON + Jinja2 模板 + xelatex 编译”的方式完成。系统先根据用户选择的模板风格，从四套 `.tex.jinja` 模板中选出目标模板，再将结构化简历数据注入到对应模板中，渲染出 `.tex` 源文件。与直接在代码中拼接 LaTeX 字符串相比，这种实现更适合维护多套风格模板，也便于后续围绕字段缺失、模板兼容和风格扩展做演进。",
        ),
        (
            "code_caption",
            "代码清单5-5 模板渲染核心代码",
        ),
        (
            "code",
            "latex_jinja_env = jinja2.Environment(\n"
            "    block_start_string=\"\\\\BLOCK{\",\n"
            "    block_end_string=\"}\",\n"
            "    variable_start_string=\"\\\\VAR{\",\n"
            "    variable_end_string=\"}\",\n"
            "    loader=jinja2.FileSystemLoader(template_dir),\n"
            ")\n"
            "template = latex_jinja_env.get_template(template_name)\n"
            "rendered_tex = template.render(resume=data)",
        ),
        (
            "paragraph",
            "渲染完成后，系统会对输出目录进行整理，将 `resume.cls` 和样式资源复制到目标路径，再调用 `xelatex` 执行 PDF 编译。当前版本除了生成优化版简历外，还会保留原始版简历的独立 PDF，用于结果页中的差异对比与质量分析。也就是说，LaTeX 层在现在的系统里不只是“出一份简历”，它还是后续评估链路的一部分。",
        ),
        (
            "code_caption",
            "代码清单5-6 PDF编译核心代码",
        ),
        (
            "code",
            "cmd = [\"xelatex\", \"-interaction=nonstopmode\", \"-halt-on-error\", tex_filename]\n"
            "result = subprocess.run(\n"
            "    cmd,\n"
            "    cwd=output_dir,\n"
            "    capture_output=True,\n"
            "    text=True,\n"
            "    encoding=\"utf-8\",\n"
            "    errors=\"ignore\",\n"
            ")\n"
            "if result.returncode != 0:\n"
            "    raise RuntimeError(\"xelatex 编译失败\")",
        ),
        (
            "paragraph",
            "考虑到真实环境中的 LaTeX 依赖并不总是完备，系统对这一步也做了较多容错处理：一方面在页面层统一将编译失败转译为更容易理解的中文错误提示；另一方面保留了 `.tex` 渲染测试能力，使模板字段不完整时可以先在不编译 PDF 的情况下验证模板输出是否可用。对于毕业设计项目来说，这种“先保底可解释，再追求完全成功”的策略，比单纯依赖本地编译环境更可靠。",
        ),
    ],
    "5.5 结果分析与可视化实现": [
        (
            "paragraph",
            "结果分析模块是系统从“能生成”走向“能解释”的关键部分。当前版本结果页不再只给出下载链接，而是围绕简历预览、质量分析、匹配度分析、词云分析、改进建议、面试问答和简历对比等多个标签页组织展示。页面在进入时优先读取会话中的 `generated_result`；若其中已经缓存分析结果，则直接复用，避免重复计算；若缓存不存在，再调用分析器完成指标计算和图表生成。这种做法兼顾了页面响应速度与分析结果的一致性。",
        ),
        (
            "code_caption",
            "代码清单5-7 结果页标签与质量分析核心代码",
        ),
        (
            "code",
            "tab_preview, tab_quality, tab_match, tab_wordcloud, tab_advice, tab_interview, tab_compare = st.tabs([\n"
            "    \"简历预览\", \"质量分析\", \"匹配度分析\", \"词云分析\", \"改进建议\", \"面试问答\", \"简历对比\"\n"
            "])\n"
            "\n"
            "quality_data = {\n"
            "    \"overlap_user\": overlap_user,\n"
            "    \"overlap_job\": overlap_job,\n"
            "    \"overlap_match\": overlap_match,\n"
            "    \"cosine_user\": cosine_user,\n"
            "    \"cosine_job\": cosine_job,\n"
            "    \"cosine_match\": cosine_match,\n"
            "}",
        ),
        (
            "paragraph",
            "在分析内容上，当前实现较前期版本更加完整。质量分析除了展示原始匹配度和优化后匹配度外，还加入了术语解释和升降对比；匹配度分析采用雷达图与评分详情并列展示；词云分析不仅给出高频词，还补充了“如何解读”的文字提示；改进建议与面试问答在模型不可用时会退化为规则生成结果。尤其是面试模块，现版本增加了 AI 面试助手，可在已有题目基础上继续追问，并对用户回答进行点评，这使结果分析不再只是静态报告，而是具备了一定交互性。",
        ),
        (
            "paragraph",
            "从实现层面看，该模块既依赖算法计算，也依赖页面编排。分析器负责输出分数、图表和建议文本，结果页组件则负责把这些结果组织成更容易理解的界面。这样的分层让后续优化变得更直接：如果需要改进算法，只要调整分析器；如果只想改善展示方式，则可以在组件层单独调整，而不必重新修改整个流程。",
        ),
    ],
    "5.6 历史记录管理实现": [
        (
            "paragraph",
            "历史记录管理模块建立在 SQLite 数据库之上，当前由 `ResumeDB` 统一封装。与项目早期只保存少量文件路径不同，现版本的 `resume_history` 表已经扩展为更完整的聚合记录结构，除了 `pdf_path`、`cv_path`、`template_style` 等结果字段，还保存了 `source_data_json`、`form_data_json`、`optimized_data_json`、`job_details_json`、`provider_name`、`model_name` 以及多类分析缓存。这样做的直接收益是：当用户从历史版本恢复时，系统不仅能重新打开旧结果，还能重新计算或直接复用质量分析、模型信息和会话状态。",
        ),
        (
            "code_caption",
            "代码清单5-8 历史记录保存核心代码",
        ),
        (
            "code",
            "c.execute(\n"
            "    \"\"\"\n"
            "    INSERT INTO resume_history\n"
            "    (user_id, source_file_path, source_data_json, original_resume_path, form_data_json,\n"
            "     optimized_data_json, pdf_path, cv_path, raw_resume_path, cv_content, template_style,\n"
            "     optimized_fields, jd_url, jd_text, job_details_json,\n"
            "     analyzer_scores, analyzer_suggestions, analyzer_questions,\n"
            "     analyzer_radar, analyzer_wordcloud, provider_name, model_name, status)\n"
            "    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'generated')\n"
            "    \"\"\",\n"
            "    values,\n"
            ")",
        ),
        (
            "paragraph",
            "在页面交互上，用户中心会根据当前用户ID读取历史记录摘要，并生成模板风格分布、模型配置分布与生成时间分布三张概览图。用户选择某条记录后，系统进一步读取完整数据，将原始简历、结构化JD、优化结果、模型配置以及分析缓存重新写回会话状态，然后自动跳转到结果分析页。也就是说，历史记录模块已经不是一个简单的“下载归档”功能，而是承担了会话恢复和多轮分析复用的职责。",
        ),
        (
            "paragraph",
            "为照顾旧版本数据库，当前实现中还加入了字段兼容机制：初始化数据库时会检查 `resume_history` 是否缺少新字段，若缺失则自动执行 `ALTER TABLE` 补齐；对于早期记录未保存 `provider_name` 和 `model_name` 的情况，管理员模块还提供了缺失元数据补全逻辑。这样的处理使论文中的系统实现不再停留在理想设计层，而是对应到了实际迭代过程中常见的数据兼容问题。",
        ),
    ],
}


CH6_UPDATES = {
    "6.1.1 测试目的": "本系统测试的目的不再局限于验证简历生成主链路是否可用，还包括检查用户中心、历史恢复、结果分析、管理员面板等扩展模块在当前版本中的协同稳定性。具体来说，一方面需要确认数据输入、模型配置、简历生成、结果分析和历史记录回溯是否满足功能预期；另一方面也需要验证管理员登录、统计看板、用户与记录管理、数据导出以及授权申请演示流程是否能够正常完成，从而保证系统在答辩展示场景下具备完整、可解释的运行闭环。",
    "6.1.2 测试工具与环境": "本系统测试采用手工测试与脚本校验结合的方式。功能测试主要通过浏览器端逐项执行页面操作，并结合项目内的最小测试集验证关键工具函数；系统运行环境为 Windows 平台、本地 Python 3.11 虚拟环境、Streamlit Web 应用与 SQLite 数据库。对于管理员模块，测试时使用系统内置的固定管理员账号 `admin` 与密码 `job-llm-admin` 进行身份校验，并重点观察统计图表渲染、用户与记录操作、CSV 导出以及授权申请流程的页面反馈是否正确。",
    "6.4.1 功能测试结果": "功能测试覆盖了数据输入、简历生成、结果分析、历史记录以及管理员面板五类能力。实际测试表明，普通用户主流程在当前版本下可以稳定完成“录入—生成—分析—保存—回溯”闭环；管理员模块则能够完成身份校验、全局统计展示、用户查看与删除、记录筛选与删除、CSV 导出以及授权申请演示等操作。测试过程中未发现会阻断答辩演示的高优先级缺陷，系统在产品完整性上较前期版本有明显提升。",
    "6.4.2 性能测试结果": "性能测试表明，系统的主要耗时仍集中在 LLM 调用、章节并发优化和 PDF 编译阶段，管理员模块并不是性能瓶颈。管理员页面的统计逻辑主要由本地 SQLite 查询和前端图表渲染构成，不涉及大模型推理，因此页面打开、筛选记录和导出 CSV 的响应时间整体较短。换言之，管理员能力的加入并没有显著拉高系统总体开销，而是以较小的额外复杂度补齐了后台运维与数据治理展示能力。",
}


def add_paragraph_before(anchor: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    paragraph = anchor.insert_paragraph_before(text)
    if style:
        paragraph.style = style
    return paragraph


def remove_paragraph(paragraph: Paragraph):
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def paragraph_has_drawing(paragraph: Paragraph) -> bool:
    xml = paragraph._element.xml
    return "<w:drawing" in xml or "<pic:pic" in xml


def is_preserved_nontext_paragraph(paragraph: Paragraph) -> bool:
    text = paragraph.text.strip()
    if paragraph_has_drawing(paragraph):
        return True
    if text.startswith("图") or text.startswith("表"):
        return True
    return False


def find_paragraph(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"未找到段落：{text}")


def find_paragraph_index(doc: Document, text: str) -> int:
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == text:
            return index
    raise ValueError(f"未找到段落：{text}")


def rewrite_section(doc: Document, heading_text: str, next_heading_text: str, blocks: list[tuple[str, str]]):
    paragraphs = doc.paragraphs
    start_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == heading_text)
    end_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == next_heading_text)
    section_paragraphs = doc.paragraphs[start_idx + 1 : end_idx]

    preserved = [p for p in section_paragraphs if is_preserved_nontext_paragraph(p)]
    anchor = preserved[0] if preserved else doc.paragraphs[end_idx]

    for paragraph in list(section_paragraphs):
        if not is_preserved_nontext_paragraph(paragraph):
            remove_paragraph(paragraph)

    for block_type, payload in blocks:
        if block_type == "paragraph":
            p = add_paragraph_before(anchor, payload, style="Normal")
            p.paragraph_format.first_line_indent = Pt(21)
            p.paragraph_format.line_spacing = 1.5
        elif block_type == "code_caption":
            p = add_paragraph_before(anchor, payload, style="Normal")
            p.paragraph_format.line_spacing = 1.5
        elif block_type == "code":
            p = add_paragraph_before(anchor, payload, style="HTML Preformatted")
        elif block_type == "blank":
            add_paragraph_before(anchor, "", style="Normal")


def insert_section_57(doc: Document):
    anchor = find_paragraph(doc, "第六章 系统测试")
    blocks = [
        ("heading", "5.7 管理员模块实现"),
        (
            "paragraph",
            "管理员模块是当前版本相较前期实现新增的一层后台能力，其目标不是构建完整权限系统，而是在现有单体应用内补齐系统级统计展示、后台维护和演示性授权流程。页面入口位于侧边导航的“管理员面板”步骤中，未认证状态下首先展示管理员登录页和申请管理员权限页。登录校验采用预置账号口令 `admin / job-llm-admin`，一旦认证通过，即在会话状态中写入 `admin_authenticated=True`，随后开放用户管理、记录管理、数据导出和授权申请等标签页。",
        ),
        (
            "paragraph",
            "在统计展示方面，管理员面板会读取数据库中的全量用户与全量简历记录，计算用户总数、今日活跃用户、累计生成记录和今日新增记录等指标，并进一步绘制近14天生成趋势、模板风格分布、模型配置分布以及用户记录数Top10等图表。由于该模块主要依赖 SQLite 查询和前端图表绘制，不涉及大模型推理，因此在交互上比生成链路更轻，适合在答辩时展示系统的后台观察能力。",
        ),
        ("code_caption", "代码清单5-9 管理员认证与统计读取核心代码"),
        (
            "code",
            "if username == ADMIN_USERNAME and password == ADMIN_PASSWORD:\n"
            "    st.session_state.admin_authenticated = True\n"
            "    st.success(\"管理员身份校验通过。\")\n"
            "    st.rerun()\n"
            "\n"
            "users_df = pd.DataFrame(db.get_all_users_overview())\n"
            "db.backfill_missing_model_metadata()\n"
            "records_df = pd.DataFrame(db.get_all_resume_records())",
        ),
        (
            "paragraph",
            "在管理操作方面，当前实现支持查看用户列表、修改用户名、删除用户及其全部历史记录、按提供商筛选生成记录、删除单条记录，以及导出用户统计与记录统计 CSV。另一个较有展示性的设计是“申请管理员权限”演示流程：普通用户可以在前台填写申请信息，管理员在后台标签页中可以执行授权通过或驳回操作。该流程目前不写入正式权限表，而是以会话态假实现方式完成，目的是在不引入复杂权限系统的前提下，完整表达“申请—审核—授权”这一后台业务链路。",
        ),
        ("blank", ""),
    ]

    for block_type, payload in blocks:
        if block_type == "heading":
            p = add_paragraph_before(anchor, payload, style="Heading 2")
            for run in p.runs:
                run.font.size = Pt(16)
        elif block_type == "paragraph":
            p = add_paragraph_before(anchor, payload, style="Normal")
            p.paragraph_format.first_line_indent = Pt(21)
            p.paragraph_format.line_spacing = 1.5
        elif block_type == "code_caption":
            p = add_paragraph_before(anchor, payload, style="Normal")
            p.paragraph_format.line_spacing = 1.5
        elif block_type == "code":
            add_paragraph_before(anchor, payload, style="HTML Preformatted")
        elif block_type == "blank":
            add_paragraph_before(anchor, "", style="Normal")


def update_ch6(doc: Document):
    for title, new_text in CH6_UPDATES.items():
        index = find_paragraph_index(doc, title)
        next_para = doc.paragraphs[index + 1]
        next_para.text = new_text

    anchor = find_paragraph(doc, "6.3 性能测试")
    blocks = [
        ("subtitle", "6.2.5 管理员模块功能测试"),
        (
            "paragraph",
            "管理员模块功能测试重点检查四类能力：一是固定账号密码登录后的身份校验是否正确；二是全局统计卡片、趋势图、模板风格分布图和模型配置分布图是否能够正常渲染；三是用户管理、记录筛选、删除操作和 CSV 导出是否能得到明确反馈；四是“申请管理员权限—管理员审批”的演示流程是否能够完整跑通。测试结果表明，该模块在当前单机部署条件下能够稳定工作，且与普通用户主流程之间不存在明显状态冲突。",
        ),
        ("blank", ""),
    ]
    for block_type, payload in blocks:
        if block_type == "subtitle":
            p = add_paragraph_before(anchor, payload, style="Normal")
            for run in p.runs:
                run.bold = True
        elif block_type == "paragraph":
            p = add_paragraph_before(anchor, payload, style="Normal")
            p.paragraph_format.first_line_indent = Pt(21)
            p.paragraph_format.line_spacing = 1.5
        elif block_type == "blank":
            add_paragraph_before(anchor, "", style="Normal")


def main():
    if not THESIS_PATH.exists():
        raise FileNotFoundError(f"未找到论文文件：{THESIS_PATH}")

    if not BACKUP_PATH.exists():
        shutil.copy2(THESIS_PATH, BACKUP_PATH)

    doc = Document(str(THESIS_PATH))

    rewrite_section(doc, "5.1 系统实现环境", "5.2 用户管理与数据输入实现", CH5_REWRITES["5.1 系统实现环境"])
    rewrite_section(doc, "5.2 用户管理与数据输入实现", "5.3 模型配置与简历生成核心实现", CH5_REWRITES["5.2 用户管理与数据输入实现"])
    rewrite_section(doc, "5.4 LaTeX模板渲染与PDF生成实现", "5.5 结果分析与可视化实现", CH5_REWRITES["5.4 LaTeX模板渲染与PDF生成实现"])
    rewrite_section(doc, "5.5 结果分析与可视化实现", "5.6 历史记录管理实现", CH5_REWRITES["5.5 结果分析与可视化实现"])
    rewrite_section(doc, "5.6 历史记录管理实现", "第六章 系统测试", CH5_REWRITES["5.6 历史记录管理实现"])
    insert_section_57(doc)
    update_ch6(doc)

    doc.save(str(THESIS_PATH))
    print(f"Updated: {THESIS_PATH}")
    print(f"Backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
